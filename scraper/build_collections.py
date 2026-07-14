"""
Genera un xlsx standalone de Smart Collections listo para importar con Matrixify.
Toma las definiciones de la hoja 'Collections (Matrixify)' del archivo del usuario
y las escribe en una hoja limpia (headers en fila 1, sin filas de notas).
"""
import os
import re
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = "/root/.claude/uploads/f4677800-7e33-5d4e-ab81-cb54e5e95596/22759230-chestershopifycolecciones_1.xlsx"
OUT = os.path.join(BASE, "chester_colecciones_matrixify.xlsx")

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(v):
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def main():
    wb = openpyxl.load_workbook(SRC)
    src = wb["Collections (Matrixify)"]
    # localizar fila de headers (contiene 'Handle')
    hrow = None
    for r in range(1, 8):
        if src.cell(r, 1).value == "Handle":
            hrow = r
            break
    headers = [src.cell(hrow, c).value for c in range(1, src.max_column + 1)]

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Collections"
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)

    n = 0
    for r in range(hrow + 1, src.max_row + 1):
        if not src.cell(r, 1).value:
            continue
        n += 1
        for c in range(1, len(headers) + 1):
            ws.cell(n + 1, c, clean(src.cell(r, c).value))

    out.save(OUT)
    print(f"guardado: {OUT}")
    print(f"colecciones: {n}")


if __name__ == "__main__":
    main()
