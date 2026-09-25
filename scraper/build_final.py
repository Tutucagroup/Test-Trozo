"""
Arma el archivo Matrixify de ACTUALIZACIÓN con:
  - Tags nuevos (recategorización multi-colección)  -> scratchpad/recat.jsonl
  - Bullets por producto (metafield custom.bullets)  -> scratchpad/bullets_out/*.jsonl

Archivo liviano (solo Handle, Command, SKU, Tags, metafield bullets) para importar
con Matrixify sin tocar precios/imágenes. Command=MERGE.
"""
import os
import re
import glob
import json
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SC = "/tmp/claude-0/-home-user-Test-Trozo/f4677800-7e33-5d4e-ab81-cb54e5e95596/scratchpad"
OUT = os.path.join(BASE, "chester_recategorizacion_bullets.xlsx")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(v):
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def main():
    # tags nuevos por handle
    recat = {}
    for line in open(os.path.join(SC, "recat.jsonl"), encoding="utf-8"):
        r = json.loads(line)
        recat[r["handle"]] = r

    # bullets por sku
    bullets = {}
    for f in sorted(glob.glob(os.path.join(SC, "bullets_out", "*.jsonl"))):
        for line in open(f, encoding="utf-8"):
            line = line.strip()
            if not line:
                continue
            try:
                b = json.loads(line)
            except Exception:
                continue
            if b.get("sku") and b.get("bullets"):
                bullets[b["sku"]] = [str(x).strip() for x in b["bullets"] if str(x).strip()][:4]
    print(f"handles con tags: {len(recat)} | skus con bullets: {len(bullets)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = ["Handle", "Command", "Variant SKU", "Tags",
               "Metafield: custom.bullets [list.single_line_text]"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)

    r = 2
    n_tags = n_bul = 0
    for handle, rec in recat.items():
        sku = rec.get("sku")
        ws.cell(r, 1, handle)
        ws.cell(r, 2, "MERGE")
        ws.cell(r, 3, sku)
        ws.cell(r, 4, clean(rec["new_tags"]))
        if rec.get("added"):
            n_tags += 1
        b = bullets.get(sku)
        if b:
            ws.cell(r, 5, clean("\n".join(b)))  # list metafield: 1 valor por línea
            n_bul += 1
        r += 1
    wb.save(OUT)
    print(f"GUARDADO {OUT} | filas: {r-1} | con tags agregados: {n_tags} | con bullets: {n_bul}")


if __name__ == "__main__":
    main()
