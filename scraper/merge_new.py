"""
Fusiona el catálogo existente (que ya funciona en Shopify) con los productos
nuevos, quitando los que fueron eliminados del sitio.

  final = existing (menos WS eliminados) + nuevos

Trabaja a nivel de bloques de producto: un bloque = fila con Command=MERGE +
sus filas de imagen siguientes (Command vacío). Conserva las hojas de
referencia (Instrucciones, Tags por categoría) del archivo existente.
"""
import os
import re
import sys
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = "Products (Matrixify)"


def load_blocks(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    ncol = ws.max_column
    blocks = []
    cur = None
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncol + 1)]
        if all(v is None for v in vals):
            continue
        cmd = ws.cell(r, 2).value  # Command
        if cmd == "MERGE":
            cur = {"rows": [vals]}
            blocks.append(cur)
        elif cur is not None:
            cur["rows"].append(vals)
    return blocks, ncol


def sku_of(block):
    # Variant SKU está en la columna 12 (índice 11)
    return block["rows"][0][11]


def main(existing, new, removed_file, out):
    removed_ws = set()
    for line in open(removed_file):
        m = re.search(r"-WS(\d+)", line)
        if m:
            removed_ws.add("WS" + m.group(1))
    print("WS eliminados:", len(removed_ws))

    ex_blocks, ncol = load_blocks(existing)
    new_blocks, _ = load_blocks(new)

    kept = [b for b in ex_blocks if sku_of(b) not in removed_ws]
    dropped = len(ex_blocks) - len(kept)
    print(f"existentes: {len(ex_blocks)} | quitados (eliminados): {dropped} | conservados: {len(kept)}")
    print(f"nuevos: {len(new_blocks)}")

    # partir del archivo existente para conservar hojas de referencia + headers/formato
    wb = openpyxl.load_workbook(existing)
    ws = wb[SHEET]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    r = 2
    for b in kept + new_blocks:
        for row in b["rows"]:
            for c, v in enumerate(row, start=1):
                if v is not None and v != "":
                    ws.cell(r, c, v)
            r += 1
    wb.save(out)
    total = len(kept) + len(new_blocks)
    print(f"GUARDADO {out} | productos: {total} | filas: {r-1}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
