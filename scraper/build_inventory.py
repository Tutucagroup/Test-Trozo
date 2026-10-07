"""
Replica el inventario del catálogo en las 6 sucursales nuevas (retiro en tienda).
Cada sucursal recibe la MISMA cantidad que el producto tiene hoy (Variant Inventory Qty).

Salida: chester_inventario_sucursales.xlsx  (import Matrixify, hoja Products)
Columnas: Handle, Variant SKU, Command=MERGE + "Inventory Available: <Sucursal>" x6.
No toca el inventario de la ubicación principal (no se incluye esa columna).
"""
import os
import re
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "chester_productos_matrixify.xlsx")
OUT = os.path.join(BASE, "chester_inventario_sucursales.xlsx")

LOCATIONS = [
    "CHESTER JUMBO - PORTAL LOS ANDES",
    "CHESTER MENDOZA PLAZA SHOPPING",
    "CHESTER PLAZA GODOY CRUZ",
    "CHESTER PLAZOLETA BARRAQUERO",
    "CHESTER PERITO MORENO",
    "CHESTER PALMARES",
]


def norm_handle(h):
    return re.sub(r"-{2,}", "-", h or "").strip("-")


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb["Products (Matrixify)"]
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    H = {h: i for i, h in enumerate(hdr)}

    out = openpyxl.Workbook()
    o = out.active
    o.title = "Products"
    headers = ["Handle", "Variant SKU", "Command"] + [f"Inventory Available: {loc}" for loc in LOCATIONS]
    for c, h in enumerate(headers, 1):
        o.cell(1, c, h)

    r = 2
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[H["Command"]] != "MERGE":
            continue
        handle = norm_handle(row[H["Handle"]])
        sku = row[H["Variant SKU"]]
        qty = row[H["Variant Inventory Qty"]]
        try:
            qty = int(float(qty))
        except (TypeError, ValueError):
            qty = 0
        o.cell(r, 1, handle)
        o.cell(r, 2, sku)
        o.cell(r, 3, "MERGE")
        for i, loc in enumerate(LOCATIONS):
            o.cell(r, 4 + i, qty)
        r += 1
        n += 1
    out.save(OUT)
    print(f"GUARDADO {OUT} | productos: {n} | sucursales: {len(LOCATIONS)}")
    print("columnas:", headers)


if __name__ == "__main__":
    main()
