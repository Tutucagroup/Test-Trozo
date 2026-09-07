"""
Carga el árbol de categorías del usuario (hoja 'Árbol & Tags' + 'Tags por categoría')
y resuelve el path completo de una categoría del sitio -> handle / tags / type.

El sitio tiene hasta 4 niveles; el árbol del usuario tiene 3. Se camina el árbol
nivel por nivel (departamento › categoría › subcategoría) con matching tolerante
(exacto / prefijo / plural) y se devuelve el nodo más profundo alcanzado.
"""
import re
import html as htmllib
import unicodedata
import openpyxl

import os as _os
_BASE = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
# El árbol (dept/cat/subcat/handle/tags/type) vive en la hoja 'Tags por categoría'
# de la plantilla committeada en el repo (equivale a 'Árbol & Tags' + Type).
TREE_XLSX = _os.path.join(_BASE, "plantilla_nuevo_producto.xlsx")
TREE_SHEET = "Tags por categoría"


def norm(s):
    if not s:
        return ""
    s = htmllib.unescape(str(s))
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _name_eq(a, b):
    """Igualdad tolerante de nombres de categoría."""
    a, b = norm(a), norm(b)
    if not a or not b:
        return False
    if a == b:
        return True
    # tolerar plural/singular y sufijos menores en el último token
    if a.rstrip("s") == b.rstrip("s"):
        return True
    if a.startswith(b) or b.startswith(a):
        return True
    return False


def load_tree():
    """Árbol anidado: {dept_name: {"node":.., "children": {cat_name: {"node":.., "children": {sub: {"node":..}}}}}}
    Lee de la hoja 'Tags por categoría': Nivel, Depto, Categoría, Subcat, Handle, Type, Tags."""
    wb = openpyxl.load_workbook(TREE_XLSX)
    ws = wb[TREE_SHEET]

    tree = {}
    for r in range(5, ws.max_row + 1):
        lvl, dep, cat, sub, handle, typ, tags = [ws.cell(r, c).value for c in range(1, 8)]
        if not handle:
            continue
        node = {
            "handle": handle,
            "tags": tags,
            "type": typ or (sub or cat or dep),
            "level": lvl,
            "names": (dep, cat, sub),
        }
        dnode = tree.setdefault(dep, {"node": None, "children": {}})
        if not cat:
            dnode["node"] = node
            continue
        cnode = dnode["children"].setdefault(cat, {"node": None, "children": {}})
        if not sub:
            cnode["node"] = node
            continue
        cnode["children"][sub] = {"node": node, "children": {}}
    return tree


class TreeMatcher:
    def __init__(self):
        self.tree = load_tree()

    def _find_child(self, children, name):
        for key, val in children.items():
            if _name_eq(key, name):
                return val
        return None

    def match_path(self, path_topdown):
        """path_topdown: [departamento, categoria, subcat, subsubcat...] -> nodo más profundo."""
        parts = [p for p in path_topdown if p]
        if not parts:
            return None
        # departamento
        dnode = self._find_child(self.tree, parts[0])
        if not dnode:
            return None
        best = dnode["node"]
        cur = dnode
        for p in parts[1:]:
            nxt = self._find_child(cur["children"], p)
            if not nxt:
                break
            if nxt["node"]:
                best = nxt["node"]
            cur = nxt
        return best
